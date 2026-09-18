#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export the risc_v verification plan to a single .xlsx workbook.

Reads the same master requirement table used by build_plan.py, so the XLSX
content is guaranteed identical to the generated HTML sheets.
Sheets + columns mirror the HTML template 1:1.
"""
import importlib.util, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("bp", os.path.join(HERE, "build_plan.py"))
bp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bp)
R, TESTS, TITLE = bp.R, bp.TESTS, bp.TITLE

# --- styles (mirroring the template look) ---
YELLOW = PatternFill("solid", fgColor="FFFF00")
BLUE   = PatternFill("solid", fgColor="8DB3E2")
THIN   = Border(*[Side(style="thin", color="B0B0B0")] * 4)
F_TITLE  = Font(name="Arial", size=10, bold=True)
F_HDR    = Font(name="Arial", size=10)
F_HDR_B  = Font(name="Arial", size=16, bold=True)   # template s5 headers ("component name" etc.)
F_DATA   = Font(name="Arial", size=11)
A_C = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
A_L = Alignment(horizontal="left",  vertical="top",    wrap_text=True)

def build_sheet(wb, name, headers, header_bold_last, widths_px, rows, first_col_centered=False):
    ws = wb.create_sheet(name)
    # column widths: px -> excel width units (~ 7 px per unit)
    for i, px in enumerate(widths_px, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8.43, round(px / 7, 1))
    ncols = len(widths_px)
    # row 1: yellow banner, merged over the used columns like the template (colspan=3)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(3, ncols))
    c = ws.cell(row=1, column=1, value=TITLE)
    c.fill, c.font, c.alignment = YELLOW, F_TITLE, Alignment(horizontal="center", vertical="bottom")
    for j in range(1, min(3, ncols) + 1):
        ws.cell(row=1, column=j).fill = YELLOW
    ws.row_dimensions[1].height = 17
    # row 2: blue headers
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=j, value=h if h is not None else "")
        c.fill = BLUE
        c.font = F_HDR_B if (header_bold_last and j == header_bold_last) else F_HDR
        c.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
        c.border = THIN
    ws.row_dimensions[2].height = 16
    # data rows
    for i, row in enumerate(rows, start=3):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = F_DATA
            c.alignment = Alignment(horizontal="center" if (first_col_centered and j == 1) else "left",
                                    vertical="top", wrap_text=True)
    ws.freeze_panes = "A3"
    return ws

wb = Workbook()
wb.remove(wb.active)

S = [(r["id"], r["desc"], r.get("note", ""), r.get("comp", "")) for r in R]
build_sheet(wb, "System",
            ["ID", "Design Requirement Description", None, "component name"],
            4, [145, 853, 738, 522], S)

G = [(r["id"], r["desc"], r["gen"], r.get("gen_obj", "")) for r in R if r.get("gen")]
build_sheet(wb, "Generation_",
            ["ID", "Design Requirement Description", "Generation", "component/Object name"],
            4, [145, 853, 701, 442], G)

C = [(r["id"], r["desc"], r["chk"], r.get("chk_prop", ""), r.get("chk_cmt", "")) for r in R if r.get("chk")]
build_sheet(wb, "Checking_",
            ["ID", "Design Requirement Description", "check", "property name", "comment"],
            4, [145, 853, 701, 184, 145], C)

V = [(r["id"], r["desc"], r["cov"], r.get("cov_cg", "")) for r in R if r.get("cov")]
build_sheet(wb, "Coverage_",
            ["ID", "Design Requirement Description", "Coverage", "covergroup"],
            4, [145, 853, 701, 470], V)

build_sheet(wb, "test list",
            ["test name", "sequance run", "stimuls generated"],
            None, [230, 176, 692], [(t[0], t[1], t[2]) for t in TESTS], first_col_centered=False)

out = os.path.join(HERE, "risc_v_verification_plan.xlsx")
wb.save(out)
print("wrote", out)
for ws in wb.worksheets:
    print(f"  sheet '{ws.title}': {ws.max_row - 2} data rows x {ws.max_column} cols")
